#!/usr/bin/env python3
"""一图全解 PPT 内容替换脚本"""

from pptx import Presentation
from pptx.util import Emu, Inches, Pt
import sys, os, json

MIN_TABLE_ROW_HEIGHT = Inches(0.25)

def replace_text(shape, new_text):
    tf = shape.text_frame
    for p in tf.paragraphs:
        for r in p.runs:
            r.text = ''
    if tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = new_text
    else:
        run = tf.paragraphs[0].add_run()
        run.text = new_text

def replace_text_keep_lines(shape, new_text):
    tf = shape.text_frame
    lines = new_text.split('\n')
    for p in tf.paragraphs:
        for r in p.runs:
            r.text = ''
    for i, line in enumerate(lines):
        if i < len(tf.paragraphs):
            if tf.paragraphs[i].runs:
                tf.paragraphs[i].runs[0].text = line
            else:
                run = tf.paragraphs[i].add_run()
                run.text = line

def replace_cell(cell, new_text):
    tf = cell.text_frame
    for p in tf.paragraphs:
        for r in p.runs:
            r.text = ''
    if tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = new_text
    else:
        run = tf.paragraphs[0].add_run()
        run.text = new_text

def walk_shapes(shapes, callback):
    for shape in shapes:
        callback(shape)
        if shape.shape_type == 6:
            walk_shapes(shape.shapes, callback)

def content_row_height(row, table):
    """Estimate row height based on longest cell content."""
    max_lines = 1
    for cell in row.cells:
        lines = cell.text_frame.text.count('\n') + 1
        max_lines = max(max_lines, lines)
    return Inches(0.25 * max_lines)

def validate_replacement(input_path, output_path, content_map, table_data=None):
    """Pre-flight validation: check all mapped shapes exist and text fits."""
    prs = Presentation(input_path)
    slide = prs.slides[0]
    found = set()
    overflows = []
    missing = []
    
    def check(shape):
        if shape.name in content_map:
            found.add(shape.name)
            new_text = content_map[shape.name]
            original_len = len(shape.text_frame.text)
            if len(new_text) > original_len * 1.1:
                overflows.append({
                    'shape': shape.name,
                    'original_len': original_len,
                    'new_len': len(new_text),
                    'overflow_ratio': len(new_text) / max(original_len, 1)
                })
    
    walk_shapes(slide.shapes, check)
    
    for name in content_map:
        if name not in found:
            missing.append(name)
    
    report = {
        'ok': len(missing) == 0 and len(overflows) == 0,
        'total_shapes': len(content_map),
        'found': len(found),
        'missing': missing,
        'overflows': overflows,
        'output_path': output_path
    }
    return report

def run_replacement(input_path, output_path, content_map, table_data=None):
    """Execute replacement with pre-flight validation."""
    validation = validate_replacement(input_path, output_path, content_map, table_data)
    if not validation['ok']:
        print(json.dumps(validation, ensure_ascii=False, indent=2))
        if validation['missing']:
            raise ValueError(f"Missing shapes: {validation['missing']}")
    
    prs = Presentation(input_path)
    slide = prs.slides[0]
    
    def process(shape):
        if shape.name in content_map:
            new_text = content_map[shape.name]
            original_len = len(shape.text_frame.text)
            if len(new_text) > original_len * 1.1:
                print(f"⚠️ {shape.name}: 替换文本({len(new_text)}字)超过原文({original_len}字)10%，可能有溢出风险")
            if '\n' in new_text:
                replace_text_keep_lines(shape, new_text)
            else:
                replace_text(shape, new_text)
            print(f"✅ {shape.name} ({len(new_text)}字)")
        
        if shape.has_table and table_data:
            table = shape.table
            for key, text in table_data.items():
                parts = key.split('_')
                row = int(parts[1])
                col = int(parts[2])
                replace_cell(table.cell(row, col), text)
            for row in table.rows:
                row.height = max(MIN_TABLE_ROW_HEIGHT, content_row_height(row, table))
            print(f"✅ {shape.name} 表格已替换 + 行高自适应")
    
    walk_shapes(slide.shapes, process)
    prs.save(output_path)
    print(f"🎉 完成：{output_path} ({os.path.getsize(output_path):,} bytes)")

if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='一图全解 PPT 替换')
    parser.add_argument('--input', required=True, help='输入 PPTX 路径')
    parser.add_argument('--output', required=True, help='输出 PPTX 路径')
    parser.add_argument('--content', required=True, help='内容映射 JSON 文件路径')
    parser.add_argument('--table', help='表格数据 JSON 文件路径')
    parser.add_argument('--dry-run', action='store_true', help='仅校验不输出')
    args = parser.parse_args()
    
    with open(args.content) as f:
        content_map = json.load(f)
    
    table_data = None
    if args.table:
        with open(args.table) as f:
            table_data = json.load(f)
    
    if args.dry_run:
        report = validate_replacement(args.input, args.output, content_map, table_data)
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        run_replacement(args.input, args.output, content_map, table_data)
